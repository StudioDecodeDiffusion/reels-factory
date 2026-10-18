"""Agent 3 : classeur Excel de suivi fin du compte, régénéré à chaque relevé et enregistré dans GitHub
(suivi/Decode_suivi_<compte>.xlsx). Source : state/suivi_<compte>.json (relevés quotidiens de l'API Instagram)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parent.parent / "suivi"
POLICE = "Arial"
ENTETE = PatternFill("solid", fgColor="111111")
TRAIT = Border(bottom=Side(style="thin", color="DDDDDD"))


def _feuille(wb, titre, entetes, largeurs, premiere=False):
    ws = wb.active if premiere else wb.create_sheet()
    ws.title = titre
    ws.append(entetes)
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    return ws


def _finir(ws, formats: dict[int, str]):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=POLICE)
            c.border = TRAIT
            if c.column in formats:
                c.number_format = formats[c.column]


def _courbe(ws, titre, cols, n, ancre, typ="ligne", y="", couleurs=()):
    ch = LineChart() if typ == "ligne" else BarChart()
    if typ == "barres_h":
        ch.type = "bar"
    ch.title, ch.height, ch.width = titre, 7.5, 16
    ch.y_axis.title = y
    ch.legend.position = "b"
    for col in cols:
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1 if typ != "barres_h" else 2, min_row=2, max_row=n))
    for s, coul in zip(ch.series, couleurs):
        if typ == "ligne":
            s.graphicalProperties.line.solidFill = coul
            s.graphicalProperties.line.width = 28000
            s.smooth = False
            s.marker.symbol = "circle"
            s.marker.size = 7
            s.marker.graphicalProperties.solidFill = coul
        else:
            s.graphicalProperties.solidFill = coul
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, ancre)


def construire(pid: str, suivi: dict, nom_compte: str) -> Path:
    wb = Workbook()
    jours = sorted(suivi.get("jours", {}))

    # 1. Jour par jour
    ws = _feuille(wb, "Jour par jour", ["Jour", "Abonnés", "Abonnés gagnés", "Vues cumulées des vidéos", "Vues du jour (vidéos)",
                                         "Portée du compte", "Vues du compte", "Visites du profil", "Comptes engagés",
                                         "Interactions", "Clics lien bio", "Vues des stories du jour", "Vidéos publiées"],
                  [12, 10, 11, 13, 12, 11, 11, 10, 11, 11, 10, 12, 10], premiere=True)
    for i, j in enumerate(jours, start=2):
        c = suivi["jours"][j]["compte"]
        vids = suivi["jours"][j]["videos"].values()
        stories = sum((s.get("views") or 0) for s in suivi.get("stories", {}).values() if (s.get("date") or "")[:10] == j)
        publiees = sum(1 for v in vids if (v.get("date") or "")[:10] == j)
        ws.append([j, c.get("abonnes"), f"=IF(ROW()=2,\"\",B{i}-B{i-1})", sum((v.get("views") or 0) for v in vids),
                   f"=IF(ROW()=2,D{i},D{i}-D{i-1})", c.get("reach"), c.get("views"), c.get("profile_views"),
                   c.get("accounts_engaged"), c.get("total_interactions"), c.get("website_clicks"), stories, publiees])
    n1 = len(jours) + 1
    _finir(ws, {k: "#,##0" for k in range(2, 14)})
    if jours:
        _courbe(ws, "Abonnés", [2], n1, "O2", y="abonnés", couleurs=["E5133A"])
        _courbe(ws, "Vues par jour : vidéos et stories", [5, 12], n1, "O18", typ="col", couleurs=["3D5AFE", "FF8A1F"])
        _courbe(ws, "Portée du compte et visites du profil", [6, 8], n1, "O34", couleurs=["8E5CFF", "00B8D4"])
        _courbe(ws, "Abonnés gagnés par jour", [3], n1, "O50", typ="col", couleurs=["2E7D32"])

    # 2. Vidéos (dernier relevé)
    ws = _feuille(wb, "Vidéos", ["Publiée le", "Titre", "Lien", "Vues", "Comptes touchés", "Likes", "Commentaires",
                                  "Partages", "Enregistrements", "Interactions", "Taux d'engagement", "Vues / comptes touchés"],
                  [16, 46, 12, 9, 11, 8, 12, 9, 13, 11, 12, 12])
    der = list(suivi["jours"][jours[-1]]["videos"].values()) if jours else []
    der.sort(key=lambda v: v.get("date") or "")
    for i, v in enumerate(der, start=2):
        ws.append([(v.get("date") or "")[:16].replace("T", " "), (v.get("legende") or "").replace("\n", " ")[:45],
                   v.get("lien"), v.get("views"), v.get("reach"), v.get("likes"), v.get("comments"), v.get("shares"),
                   v.get("saved"), v.get("total_interactions"),
                   f"=IFERROR((F{i}+G{i}+H{i}+I{i})/IF(E{i}>0,E{i},D{i}),0)", f"=IFERROR(D{i}/E{i},0)"])
        if v.get("lien"):
            ws.cell(row=i, column=3).hyperlink = v["lien"]
            ws.cell(row=i, column=3).value = "ouvrir"
    n2 = len(der) + 1
    _finir(ws, {**{k: "#,##0" for k in range(4, 11)}, 11: "0.0%", 12: "0.00"})
    for i in range(2, n2 + 1):
        ws.cell(row=i, column=3).font = Font(name=POLICE, color="0563C1", underline="single")
    if der:
        _courbe(ws, "Vues par vidéo", [4], n2, f"A{n2 + 3}", typ="barres_h", couleurs=["3D5AFE"])
        _courbe(ws, "Likes, commentaires, partages, enregistrements", [6, 7, 8, 9], n2, f"F{n2 + 3}", typ="barres_h",
                couleurs=["FF4F81", "8E5CFF", "2E7D32", "FF8A1F"])
        _courbe(ws, "Taux d'engagement", [11], n2, f"A{n2 + 19}", typ="barres_h", couleurs=["E5133A"])

    # 3. Évolution des vues de chaque vidéo (une colonne par vidéo, une ligne par jour de relevé)
    ws = _feuille(wb, "Évolution par vidéo", ["Jour"] + [(v.get("legende") or "")[:28].replace("\n", " ") for v in der],
                  [12] + [16] * len(der))
    ids = [k for k, _ in sorted(suivi["jours"][jours[-1]]["videos"].items(), key=lambda kv: kv[1].get("date") or "")] if jours else []
    for j in jours:
        vids = suivi["jours"][j]["videos"]
        ws.append([j] + [(vids.get(k) or {}).get("views") for k in ids])
    _finir(ws, {k: "#,##0" for k in range(2, len(ids) + 2)})
    if ids and jours:
        _courbe(ws, "Vues cumulées de chaque vidéo, jour après jour", list(range(2, len(ids) + 2)), len(jours) + 1,
                f"A{len(jours) + 4}")

    # 4. Stories
    ws = _feuille(wb, "Stories", ["Publiée le", "Vues", "Comptes touchés", "Réponses", "Partages", "Abonnements",
                                   "Visites du profil", "Interactions", "Lien"], [16, 8, 11, 9, 9, 11, 11, 11, 10])
    st = sorted(suivi.get("stories", {}).values(), key=lambda s: s.get("date") or "")
    for i, s in enumerate(st, start=2):
        ws.append([(s.get("date") or "")[:16].replace("T", " "), s.get("views"), s.get("reach"), s.get("replies"),
                   s.get("shares"), s.get("follows"), s.get("profile_visits"), s.get("total_interactions"), "ouvrir"])
        if s.get("lien"):
            ws.cell(row=i, column=9).hyperlink = s["lien"]
    _finir(ws, {k: "#,##0" for k in range(2, 9)})
    if st:
        _courbe(ws, "Vues des stories", [2], len(st) + 1, "K2", typ="col", couleurs=["FF8A1F"])

    # 5. Lisez-moi
    ws = wb.create_sheet("À lire")
    for ligne in [f"Décodé · suivi du compte {nom_compte}",
                  "Mis à jour automatiquement par l'agent 3 à chaque relevé (9 h et 16 h, heure de Paris).",
                  "Source : API Instagram (insights). Les vues et partages sont cumulés depuis la publication.",
                  "« Abonnés gagnés » et « Vues du jour » = différence avec le relevé de la veille (formules).",
                  "Taux d'engagement = (likes + commentaires + partages + enregistrements) / comptes touchés.",
                  "Les statistiques d'une story ne sont lisibles que pendant ses 24 h : dernier relevé conservé."]:
        ws.append([ligne])
    ws["A1"].font = Font(name=POLICE, bold=True, size=14)
    for r in range(2, 7):
        ws.cell(row=r, column=1).font = Font(name=POLICE)
    ws.column_dimensions["A"].width = 100

    wb.calculation.fullCalcOnLoad = True
    RACINE.mkdir(exist_ok=True)
    out = RACINE / f"Decode_suivi_{pid}.xlsx"
    wb.save(out)
    return out
